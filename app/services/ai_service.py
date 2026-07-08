"""
AI Service — Abstraction layer cho Gemini và Groq.
Chọn provider qua biến môi trường AI_PROVIDER=gemini|groq.
"""
import os
import json
import re
from dotenv import load_dotenv

load_dotenv()

AI_PROVIDER = os.getenv("AI_PROVIDER", "gemini").lower()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
GROQ_API_KEY = os.getenv("GROQ_API_KEY", "")

PARSE_PROMPT = """
Bạn là AI phân tích dữ liệu. Phân tích đoạn text dưới đây và trích xuất thông tin về domain và server.
Trả về JSON theo format sau (chỉ trả về JSON, không giải thích thêm):

{
  "domains": [
    {
      "domain_name": "tên domain",
      "registrar": "nhà đăng ký (GoDaddy/VNPT/...)",
      "expiry_date": "YYYY-MM-DD hoặc null",
      "auto_renew": true/false,
      "notes": "ghi chú"
    }
  ],
  "servers": [
    {
      "label": "tên/nhãn server",
      "provider": "nhà cung cấp (Vultr/DigitalOcean/...)",
      "ip_address": "IP hoặc null",
      "type": "VPS/Shared Hosting/Dedicated/Cloud",
      "expiry_date": "YYYY-MM-DD hoặc null",
      "monthly_cost": số tiền/tháng hoặc null,
      "notes": "ghi chú"
    }
  ]
}

Text cần phân tích:
"""


def extract_text_from_file(file_path: str, original_filename: str) -> str:
    """
    Đọc text từ file được upload.
    Hỗ trợ: .txt, .docx, .xlsx, .jpg, .png, .jpeg
    """
    ext = original_filename.lower().rsplit(".", 1)[-1]

    if ext == "txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    elif ext == "docx":
        try:
            from docx import Document
            doc = Document(file_path)
            return "\n".join([p.text for p in doc.paragraphs])
        except Exception as e:
            return f"Lỗi đọc docx: {e}"

    elif ext == "xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            text_parts = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join([str(c) for c in row if c is not None])
                    if row_text.strip():
                        text_parts.append(row_text)
            return "\n".join(text_parts)
        except Exception as e:
            return f"Lỗi đọc xlsx: {e}"

    elif ext in ("jpg", "jpeg", "png"):
        try:
            import pytesseract
            from PIL import Image
            img = Image.open(file_path)
            return pytesseract.image_to_string(img, lang="vie+eng")
        except Exception as e:
            return f"Lỗi OCR ảnh: {e}. Hãy cài Tesseract OCR."

    else:
        return "Định dạng file không được hỗ trợ."


def parse_with_ai(text: str) -> dict:
    """
    Gửi text lên AI để phân tích.
    Tự động chọn Gemini hoặc Groq theo AI_PROVIDER trong .env.
    """
    if AI_PROVIDER == "groq" and GROQ_API_KEY:
        return _parse_groq(text)
    elif GEMINI_API_KEY:
        return _parse_gemini(text)
    else:
        return _parse_fallback(text)


def _parse_gemini(text: str) -> dict:
    """Gọi Google Gemini API (gemini-2.0-flash — miễn phí)."""
    try:
        import google.generativeai as genai
        genai.configure(api_key=GEMINI_API_KEY)
        model = genai.GenerativeModel("gemini-2.0-flash")
        response = model.generate_content(PARSE_PROMPT + text)
        raw = response.text.strip()
        return _extract_json(raw)
    except Exception as e:
        return {"error": str(e), "domains": [], "servers": []}


def _parse_groq(text: str) -> dict:
    """Gọi Groq API (llama-3.3-70b-versatile — miễn phí)."""
    try:
        from groq import Groq
        client = Groq(api_key=GROQ_API_KEY)
        chat = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[
                {"role": "system", "content": "Bạn là AI phân tích dữ liệu, chỉ trả về JSON."},
                {"role": "user", "content": PARSE_PROMPT + text},
            ],
            temperature=0.1,
        )
        raw = chat.choices[0].message.content.strip()
        return _extract_json(raw)
    except Exception as e:
        return {"error": str(e), "domains": [], "servers": []}


def _parse_fallback(text: str) -> dict:
    """Fallback khi không có API key — trả về thông báo."""
    return {
        "error": "Chưa cấu hình API key AI. Vui lòng thêm GEMINI_API_KEY hoặc GROQ_API_KEY vào file .env",
        "domains": [],
        "servers": [],
    }


def _extract_json(raw: str) -> dict:
    """Trích xuất JSON từ response AI (có thể có text bọc ngoài)."""
    # Thử parse trực tiếp
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        pass
    # Tìm JSON block trong text
    match = re.search(r"\{.*\}", raw, re.DOTALL)
    if match:
        try:
            return json.loads(match.group())
        except json.JSONDecodeError:
            pass
    return {"error": "AI trả về format không hợp lệ", "raw": raw, "domains": [], "servers": []}
